"""
Patent Vector Search Evaluation Script
Evaluates vector and hybrid search performance for patent documents
"""
import asyncio
import logging
import os
import sys
from datetime import datetime
from typing import Dict, List, Optional

import pandas as pd
from azure.cosmos import CosmosClient
from dotenv import load_dotenv
from elasticsearch import Elasticsearch
from openai import AzureOpenAI
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Load environment variables
load_dotenv()

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)


class PatentVectorEvaluator:
    """Evaluates vector search performance for patent documents"""

    def __init__(self):
        """Initialize the evaluator with necessary clients"""
        # Cosmos DB configuration
        self.cosmos_endpoint = os.getenv("COSMOS_ENDPOINT")
        self.cosmos_key = os.getenv("COSMOS_KEY")
        self.cosmos_database = os.getenv("COSMOS_DATABASE", "patent_json")
        self.cosmos_container = os.getenv("COSMOS_CONTAINER", "patent_csv1")

        # Elasticsearch configuration
        self.es_host = os.getenv("ELASTICSEARCH_HOST", "localhost")
        self.es_port = int(os.getenv("ELASTICSEARCH_PORT", "9200"))
        self.es_index = os.getenv("ELASTICSEARCH_INDEX", "patent_vectors")

        # Azure OpenAI configuration
        self.openai_endpoint = os.getenv("AZURE_OPENAI_ENDPOINT", "https://patent-openai.openai.azure.com/")
        self.openai_api_key = os.getenv("AZURE_OPENAI_API_KEY")
        self.openai_deployment = os.getenv("AZURE_OPENAI_DEPLOYMENT", "text-embedding-3-large")
        self.openai_api_version = os.getenv("AZURE_OPENAI_API_VERSION", "2024-12-01-preview")

        # Initialize clients
        self._init_cosmos_client()
        self._init_elasticsearch_client()
        self._init_openai_client()

        # Results storage
        self.results = []

    def _init_cosmos_client(self):
        """Initialize Cosmos DB client"""
        try:
            if not self.cosmos_endpoint or not self.cosmos_key:
                raise ValueError("Cosmos DB credentials not found in environment variables")

            self.cosmos_client = CosmosClient(self.cosmos_endpoint, self.cosmos_key)
            self.database = self.cosmos_client.get_database_client(self.cosmos_database)
            self.container = self.database.get_container_client(self.cosmos_container)
            logger.info(f"Connected to Cosmos DB: {self.cosmos_database}/{self.cosmos_container}")
        except Exception as e:
            logger.error(f"Failed to initialize Cosmos DB client: {e}")
            raise

    def _init_elasticsearch_client(self):
        """Initialize Elasticsearch client"""
        try:
            # Build proper URL
            es_url = f"http://{self.es_host}:{self.es_port}"
            logger.info(f"Attempting to connect to Elasticsearch at: {es_url}")

            self.es_client = Elasticsearch(
                es_url,
                verify_certs=False,
                request_timeout=30,
                max_retries=3,
                retry_on_timeout=True
            )
            # Test connection
            if not self.es_client.ping():
                raise ConnectionError("Cannot connect to Elasticsearch")
            logger.info(f"Connected to Elasticsearch at {self.es_host}:{self.es_port}")
        except Exception as e:
            logger.error(f"Failed to initialize Elasticsearch client: {e}")
            raise

    def _init_openai_client(self):
        """Initialize Azure OpenAI client"""
        try:
            if not self.openai_api_key:
                raise ValueError("Azure OpenAI API key not found in environment variables")

            # Clear proxy environment variables
            for proxy_var in ['http_proxy', 'https_proxy', 'HTTP_PROXY', 'HTTPS_PROXY']:
                if proxy_var in os.environ:
                    del os.environ[proxy_var]

            self.openai_client = AzureOpenAI(
                api_version=self.openai_api_version,
                azure_endpoint=self.openai_endpoint,
                api_key=self.openai_api_key
            )
            logger.info("Azure OpenAI client initialized successfully")
        except Exception as e:
            logger.error(f"Failed to initialize Azure OpenAI client: {e}")
            raise

    def get_patent_from_cosmos(self, patent_id: str) -> Optional[Dict]:
        """
        Get patent document from Cosmos DB by patent_id

        Args:
            patent_id: Patent ID to search for

        Returns:
            Patent document or None if not found
        """
        try:
            # Query to find patent by reference.syutugan
            query = """
            SELECT * FROM c
            WHERE ARRAY_CONTAINS(c.metadata.reference.syutugan, @patent_id)
            """

            items = list(self.container.query_items(
                query=query,
                parameters=[
                    {"name": "@patent_id", "value": patent_id}
                ],
                enable_cross_partition_query=True
            ))

            if items:
                logger.info(f"Found patent {patent_id} in Cosmos DB")
                return items[0]
            else:
                logger.warning(f"Patent {patent_id} not found in Cosmos DB")
                return None

        except Exception as e:
            logger.error(f"Error querying Cosmos DB for patent {patent_id}: {e}")
            return None

    def generate_embedding(self, text: str) -> Optional[List[float]]:
        """
        Generate embedding for text using Azure OpenAI

        Args:
            text: Text to generate embedding for

        Returns:
            Embedding vector or None if failed
        """
        try:
            response = self.openai_client.embeddings.create(
                model=self.openai_deployment,
                input=text
            )
            return response.data[0].embedding
        except Exception as e:
            logger.error(f"Error generating embedding: {e}")
            return None

    def vector_search(self, query_vector: List[float], k: int = 10) -> List[Dict]:
        """
        Perform vector search in Elasticsearch

        Args:
            query_vector: Query embedding vector
            k: Number of results to return

        Returns:
            List of search results
        """
        try:
            query = {
                "knn": {
                    "field": "summary_vector",
                    "query_vector": query_vector,
                    "k": min(k, 100),  # Limit to 100 for performance
                    "num_candidates": min(k * 10, 1000)
                }
            }

            response = self.es_client.search(
                index=self.es_index,
                body=query,
                size=min(k, 100)
            )

            results = []
            for hit in response["hits"]["hits"]:
                result = {
                    "patent_id": hit["_id"],
                    "score": hit["_score"],
                    "reference": hit["_source"].get("reference", {}),
                    "title": hit["_source"].get("title", ""),
                    "summary": hit["_source"].get("summary", "")[:200]  # First 200 chars
                }
                results.append(result)

            return results

        except Exception as e:
            logger.error(f"Error in vector search: {e}")
            return []

    def hybrid_search(self, query_text: str, query_vector: List[float], k: int = 10) -> List[Dict]:
        """
        Perform hybrid search (text + vector) in Elasticsearch

        Args:
            query_text: Query text
            query_vector: Query embedding vector
            k: Number of results to return

        Returns:
            List of search results
        """
        try:
            # Build hybrid query with both text and vector components
            query = {
                "query": {
                    "bool": {
                        "should": [
                            {
                                "multi_match": {
                                    "query": query_text,
                                    "fields": ["title^2", "summary", "keywords"],
                                    "type": "best_fields",
                                    "boost": 0.3
                                }
                            }
                        ]
                    }
                },
                "knn": {
                    "field": "summary_vector",
                    "query_vector": query_vector,
                    "k": min(k, 100),
                    "num_candidates": min(k * 10, 1000),
                    "boost": 0.7
                },
                "size": min(k, 100)
            }

            response = self.es_client.search(
                index=self.es_index,
                body=query
            )

            results = []
            for hit in response["hits"]["hits"]:
                result = {
                    "patent_id": hit["_id"],
                    "score": hit["_score"],
                    "reference": hit["_source"].get("reference", {}),
                    "title": hit["_source"].get("title", ""),
                    "summary": hit["_source"].get("summary", "")[:200]
                }
                results.append(result)

            return results

        except Exception as e:
            logger.error(f"Error in hybrid search: {e}")
            return []

    def check_match(self, search_results: List[Dict], target_patent_id: str, k: int) -> int:
        """
        Check if target patent ID appears in top-k search results

        Args:
            search_results: List of search results
            target_patent_id: Patent ID to look for
            k: Top-k results to consider

        Returns:
            1 if found in top-k, 0 otherwise
        """
        # Get top-k results
        top_k_results = search_results[:k]

        # Check each result for the target patent
        for result in top_k_results:
            # Get reference.syutugan directly from the result
            reference = result.get("reference", {})
            syutugan_list = reference.get("syutugan", [])

            # Check if target patent is in syutugan list
            if target_patent_id in syutugan_list:
                return 1

        return 0

    async def process_single_case(self, case_data: Dict) -> Dict:
        """
        Process a single test case

        Args:
            case_data: Dictionary with case_id, syutugan, ax_docs

        Returns:
            Dictionary with evaluation results
        """
        case_id = case_data["case_id"]
        syutugan = case_data["syutugan"]
        ax_docs = case_data["ax_docs"]

        logger.info(f"Processing {case_id}: syutugan={syutugan}, ax_docs={ax_docs}")

        result = {
            "case_id": case_id,
            "syutugan": syutugan,
            "ax_docs": ax_docs,
            "vector_top1_match": 0,
            "vector_top5_match": 0,
            "vector_top10_match": 0,
            "vector_top50_match": 0,
            "vector_top100_match": 0,
            "hybrid_top1_match": 0,
            "hybrid_top5_match": 0,
            "hybrid_top10_match": 0,
            "hybrid_top50_match": 0,
            "hybrid_top100_match": 0,
            "status": "success",
            "vector_hits": [],
            "hybrid_hits": []
        }

        try:
            # Step 1: Get patent from Cosmos DB
            patent_doc = self.get_patent_from_cosmos(syutugan)
            if not patent_doc:
                result["status"] = "patent_not_found"
                return result

            # Step 2: Get summary
            summary = patent_doc.get("summary", "")
            if not summary:
                result["status"] = "no_summary"
                return result

            # Step 3: Generate embedding
            embedding = self.generate_embedding(summary)
            if not embedding:
                result["status"] = "embedding_failed"
                return result

            # Step 4: Perform vector search (get top 100)
            vector_results = self.vector_search(embedding, k=100)
            if vector_results:
                result["vector_top1_match"] = self.check_match(vector_results, ax_docs, 1)
                result["vector_top5_match"] = self.check_match(vector_results, ax_docs, 5)
                result["vector_top10_match"] = self.check_match(vector_results, ax_docs, 10)
                result["vector_top50_match"] = self.check_match(vector_results, ax_docs, 50)
                result["vector_top100_match"] = self.check_match(vector_results, ax_docs, 100)

                # Store top 10 hits for detailed output
                result["vector_hits"] = [
                    {
                        "rank": i + 1,
                        "patent_id": hit["patent_id"],
                        "score": hit["score"],
                        "title": hit.get("title", "")[:100],
                        "is_match": ax_docs in hit.get("reference", {}).get("syutugan", [])
                    }
                    for i, hit in enumerate(vector_results[:10])
                ]

            # Step 5: Perform hybrid search (get top 100)
            hybrid_results = self.hybrid_search(summary[:500], embedding, k=100)
            if hybrid_results:
                result["hybrid_top1_match"] = self.check_match(hybrid_results, ax_docs, 1)
                result["hybrid_top5_match"] = self.check_match(hybrid_results, ax_docs, 5)
                result["hybrid_top10_match"] = self.check_match(hybrid_results, ax_docs, 10)
                result["hybrid_top50_match"] = self.check_match(hybrid_results, ax_docs, 50)
                result["hybrid_top100_match"] = self.check_match(hybrid_results, ax_docs, 100)

                # Store top 10 hits for detailed output
                result["hybrid_hits"] = [
                    {
                        "rank": i + 1,
                        "patent_id": hit["patent_id"],
                        "score": hit["score"],
                        "title": hit.get("title", "")[:100],
                        "is_match": ax_docs in hit.get("reference", {}).get("syutugan", [])
                    }
                    for i, hit in enumerate(hybrid_results[:10])
                ]

        except Exception as e:
            logger.error(f"Error processing {case_id}: {e}")
            result["status"] = f"error: {str(e)}"

        return result

    async def process_all_cases(self, csv_file: str):
        """
        Process all cases from CSV file

        Args:
            csv_file: Path to CSV file with test cases
        """
        # Read CSV file
        df = pd.read_csv(csv_file)
        logger.info(f"Loaded {len(df)} cases from {csv_file}")

        # Process each case
        for index, row in df.iterrows():
            case_data = {
                "case_id": row["case_id"],
                "syutugan": row["syutugan"],
                "ax_docs": row["ax_docs"]
            }

            result = await self.process_single_case(case_data)
            self.results.append(result)

            # Progress update
            if (index + 1) % 10 == 0:
                logger.info(f"Processed {index + 1}/{len(df)} cases")

        logger.info(f"Completed processing all {len(df)} cases")

    def save_to_excel(self, output_file: str):
        """
        Save results to Excel file with formatting

        Args:
            output_file: Path to output Excel file
        """
        # Prepare main results DataFrame (without hit details)
        main_results = []
        for r in self.results:
            main_result = {k: v for k, v in r.items() if k not in ['vector_hits', 'hybrid_hits']}
            main_results.append(main_result)

        df_main = pd.DataFrame(main_results)

        # Prepare vector hits details
        vector_hits_data = []
        for r in self.results:
            case_id = r['case_id']
            for hit in r.get('vector_hits', []):
                vector_hits_data.append({
                    'case_id': case_id,
                    'rank': hit['rank'],
                    'patent_id': hit['patent_id'],
                    'score': hit['score'],
                    'title': hit['title'],
                    'is_match': hit['is_match']
                })

        df_vector_hits = pd.DataFrame(vector_hits_data) if vector_hits_data else pd.DataFrame()

        # Prepare hybrid hits details
        hybrid_hits_data = []
        for r in self.results:
            case_id = r['case_id']
            for hit in r.get('hybrid_hits', []):
                hybrid_hits_data.append({
                    'case_id': case_id,
                    'rank': hit['rank'],
                    'patent_id': hit['patent_id'],
                    'score': hit['score'],
                    'title': hit['title'],
                    'is_match': hit['is_match']
                })

        df_hybrid_hits = pd.DataFrame(hybrid_hits_data) if hybrid_hits_data else pd.DataFrame()

        # Create Excel writer
        with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
            # Write main results
            df_main.to_excel(writer, sheet_name='Evaluation Results', index=False)

            # Write vector hits details
            if not df_vector_hits.empty:
                df_vector_hits.to_excel(writer, sheet_name='Vector Search Hits', index=False)

            # Write hybrid hits details
            if not df_hybrid_hits.empty:
                df_hybrid_hits.to_excel(writer, sheet_name='Hybrid Search Hits', index=False)

            # Format main results sheet
            worksheet = writer.sheets['Evaluation Results']
            self._format_worksheet(worksheet, df_main)

            # Format vector hits sheet
            if not df_vector_hits.empty:
                worksheet_vector = writer.sheets['Vector Search Hits']
                self._format_worksheet(worksheet_vector, df_vector_hits)

            # Format hybrid hits sheet
            if not df_hybrid_hits.empty:
                worksheet_hybrid = writer.sheets['Hybrid Search Hits']
                self._format_worksheet(worksheet_hybrid, df_hybrid_hits)

        logger.info(f"Results saved to {output_file}")

    def _format_worksheet(self, worksheet, df):
        """Format a worksheet with headers and column widths"""
        from openpyxl.formatting.rule import CellIsRule

        # Format header row
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        for col in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except (TypeError, AttributeError):
                    pass

            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width

        # Add conditional formatting for match columns
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        # Find columns with 'match' in the name
        for idx, col_name in enumerate(df.columns, 1):
            if 'match' in col_name.lower():
                col_letter = get_column_letter(idx)
                # Green for matches (1)
                worksheet.conditional_formatting.add(
                    f'{col_letter}2:{col_letter}{len(df) + 1}',
                    CellIsRule(operator='equal', formula=['1'], fill=green_fill)
                )
                # Red for no matches (0)
                worksheet.conditional_formatting.add(
                    f'{col_letter}2:{col_letter}{len(df) + 1}',
                    CellIsRule(operator='equal', formula=['0'], fill=red_fill)
                )
            elif col_name == 'is_match':
                col_letter = get_column_letter(idx)
                # Green for True
                worksheet.conditional_formatting.add(
                    f'{col_letter}2:{col_letter}{len(df) + 1}',
                    CellIsRule(operator='equal', formula=['TRUE'], fill=green_fill)
                )
                # Red for False
                worksheet.conditional_formatting.add(
                    f'{col_letter}2:{col_letter}{len(df) + 1}',
                    CellIsRule(operator='equal', formula=['FALSE'], fill=red_fill)
                )

    def print_summary(self):
        """Print summary statistics"""
        if not self.results:
            logger.warning("No results to summarize")
            return

        df = pd.DataFrame(self.results)

        # Filter successful results
        success_df = df[df['status'] == 'success']

        if len(success_df) == 0:
            logger.warning("No successful evaluations")
            return

        logger.info("=" * 60)
        logger.info("EVALUATION SUMMARY")
        logger.info("=" * 60)
        logger.info(f"Total cases processed: {len(df)}")
        logger.info(f"Successful evaluations: {len(success_df)}")
        logger.info(f"Failed evaluations: {len(df) - len(success_df)}")

        if len(df) > len(success_df):
            failure_reasons = df[df['status'] != 'success']['status'].value_counts()
            logger.info("\nFailure reasons:")
            for reason, count in failure_reasons.items():
                logger.info(f"  - {reason}: {count}")

        logger.info("\nMatch rates (successful evaluations only):")

        # Calculate match rates
        metrics = [
            ("Vector Search Top-1", "vector_top1_match"),
            ("Vector Search Top-5", "vector_top5_match"),
            ("Vector Search Top-10", "vector_top10_match"),
            ("Vector Search Top-50", "vector_top50_match"),
            ("Vector Search Top-100", "vector_top100_match"),
            ("Hybrid Search Top-1", "hybrid_top1_match"),
            ("Hybrid Search Top-5", "hybrid_top5_match"),
            ("Hybrid Search Top-10", "hybrid_top10_match"),
            ("Hybrid Search Top-50", "hybrid_top50_match"),
            ("Hybrid Search Top-100", "hybrid_top100_match")
        ]

        for label, column in metrics:
            matches = success_df[column].sum()
            total = len(success_df)
            rate = (matches / total * 100) if total > 0 else 0
            logger.info(f"  {label}: {matches}/{total} ({rate:.1f}%)")

        logger.info("=" * 60)


async def main():
    """Main function"""
    # Input and output files
    input_csv = "filtered_test_cases_20250813_174350.csv"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_excel = f"patent_evaluation_results_{timestamp}.xlsx"

    # Check if input file exists
    if not os.path.exists(input_csv):
        logger.error(f"Input file not found: {input_csv}")
        sys.exit(1)

    # Create evaluator instance
    evaluator = PatentVectorEvaluator()

    try:
        # Process all cases
        await evaluator.process_all_cases(input_csv)

        # Save results to Excel
        evaluator.save_to_excel(output_excel)

        # Print summary
        evaluator.print_summary()

        logger.info(f"Evaluation completed. Results saved to: {output_excel}")

    except Exception as e:
        logger.error(f"Evaluation failed: {e}")
        sys.exit(1)


if __name__ == "__main__":
    asyncio.run(main())
